"""Convert docs/REPORT.md to docs/REPORT.docx using python-docx.

Lightweight Markdown subset: H1/H2/H3, paragraphs, fenced code blocks,
GitHub-style tables, image references, and bullet lists.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "REPORT.md"
DST = ROOT / "docs" / "REPORT.docx"


def _add_code(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


_INLINE_RE = re.compile(
    r"\*\*(?P<bold>[^*]+)\*\*"
    r"|`(?P<code>[^`]+)`"
    r"|\[(?P<linktext>[^\]]+)\]\((?P<linkurl>[^)]+)\)"
    r"|<(?P<autolink>https?://[^>]+)>"
)


def _add_rich_runs(paragraph, text: str) -> None:
    """Append runs to ``paragraph`` parsing **bold**, `code`, and [text](url)."""
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group("bold") is not None:
            r = paragraph.add_run(m.group("bold"))
            r.bold = True
        elif m.group("code") is not None:
            r = paragraph.add_run(m.group("code"))
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif m.group("linktext") is not None:
            paragraph.add_run(m.group("linktext"))
            url = m.group("linkurl")
            if url and not url.startswith(("#", "REPORT.")):
                r = paragraph.add_run(f" ({url})")
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif m.group("autolink") is not None:
            r = paragraph.add_run(m.group("autolink"))
            r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _set_cell_rich(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if bold:
        # render entire cell as bold; still parse code spans for monospacing
        for chunk in re.split(r"(`[^`]+`)", text):
            if chunk.startswith("`") and chunk.endswith("`"):
                r = p.add_run(chunk[1:-1])
                r.bold = True
                r.font.name = "Consolas"
            else:
                r = p.add_run(chunk)
                r.bold = True
    else:
        _add_rich_runs(p, text)


def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    n = len(header)
    t = doc.add_table(rows=1 + len(rows), cols=n)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        _set_cell_rich(t.rows[0].cells[i], h, bold=True)
    for ri, row in enumerate(rows, start=1):
        padded = (row + [""] * n)[:n]
        for ci, val in enumerate(padded):
            _set_cell_rich(t.rows[ri].cells[ci], val)
    doc.add_paragraph()


_IMG_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")


def _add_image(doc: Document, ref: str) -> None:
    img = (SRC.parent / ref).resolve()
    if img.exists():
        try:
            doc.add_picture(str(img), width=Inches(5.5))
        except Exception:
            doc.add_paragraph(f"[image: {img.name}]")
    else:
        doc.add_paragraph(f"[image not found: {ref}]")


def convert() -> Path:
    if not SRC.exists():
        raise SystemExit(f"Missing {SRC}")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i, in_code, code_buf, table_buf = 0, False, [], []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in table_buf]
        # drop the markdown alignment row
        rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
        if rows:
            _add_table(doc, rows[0], rows[1:])
        table_buf = []

    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            if in_code:
                _add_code(doc, "\n".join(code_buf))
                code_buf, in_code = [], False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line); i += 1; continue

        if line.startswith("|") and "|" in line[1:]:
            table_buf.append(line); i += 1; continue
        else:
            flush_table()

        m = _IMG_RE.search(line)
        if m:
            _add_image(doc, m.group(1)); i += 1; continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_rich_runs(p, line[2:].strip())
        elif line.startswith("---"):
            doc.add_paragraph("")  # section break
        elif line.strip() == "":
            doc.add_paragraph("")
        else:
            p = doc.add_paragraph()
            _add_rich_runs(p, line)
        i += 1

    flush_table()
    doc.save(DST)
    return DST


def main() -> int:
    p = convert()
    print(f"Wrote {p} ({p.stat().st_size/1024:.1f} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
